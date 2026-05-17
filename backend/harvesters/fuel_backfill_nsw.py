# backend/harvesters/fuel_backfill_nsw.py
# Team: 68 | Name: Kuiyuan Shu | Student ID: 1345012
# NSW 油价历史回填：data.nsw.gov.au → Redis → official_data (ES)
# 一次性脚本，回填 2026-02 ~ 2026-04（5 月走 daily harvester）。
# Feb / Mar 走 CKAN datastore_search API，Apr 走 XLSX 下载。
# 用法： python fuel_backfill_nsw.py

import io
import json
import urllib.request
from collections import defaultdict
from datetime import datetime

import redis
from openpyxl import load_workbook

# ── Redis 配置（与 fuel_harvester.py 保持一致）────────────────
REDIS_HOST = "redis.redis.svc.cluster.local"
REDIS_PORT = 6379
REDIS_PASS = ""
QUEUE_KEY  = "official_data:queue"
DEDUP_KEY  = "official_data:fuel:seen"

# ── 数据源 ───────────────────────────────────────────────────
CKAN_API = "https://data.nsw.gov.au/data/api/action/datastore_search"
PAGE_SIZE = 32000   # CKAN datastore_search 单次最大

# Feb / Mar 用 CKAN datastore API（resource_id）
DATASTORE_RESOURCES = {
    "2026-02": "3786820f-8efd-4b13-b2ba-56096a9d42b3",
    "2026-03": "5576690a-2d64-4c61-a6a8-fa2975fc9d4e",
}

# Apr 用 XLSX 直链下载
XLSX_URLS = {
    "2026-04": ("https://data.nsw.gov.au/data/dataset/a97a46fc-2bdd-4b90-ac7f-0cb1e8d7ac3b"
                "/resource/4a9d6b0e-fcaf-4d8b-b443-6f214ca7c8ce/download/fuelcheck_pricehistory_apr2026.xlsx"),
}

# ── 油种映射（与 fuel_harvester 的 ALLOWED_TYPES 对齐）───────
FUEL_MAP = {
    "U91": "u91",
    "P95": "u95",
    "P98": "u98",
    "DL":  "diesel",
}
ALLOWED_TYPES = set(FUEL_MAP.values())


# ══════════════════════════════════════════════════════════════
# 数据源 A：CKAN datastore_search API（CSV 后台）
# ══════════════════════════════════════════════════════════════

def fetch_datastore(resource_id):
    """分页拉取一个 resource 的全部 records。返回 [dict, ...]"""
    records = []
    offset = 0
    while True:
        url = f"{CKAN_API}?resource_id={resource_id}&limit={PAGE_SIZE}&offset={offset}"
        try:
            with urllib.request.urlopen(url, timeout=120) as resp:
                payload = json.loads(resp.read().decode("utf-8"))
        except Exception as e:
            print(f"[BACKFILL-NSW] CKAN 请求失败 offset={offset}: {e}")
            break

        if not payload.get("success"):
            print(f"[BACKFILL-NSW] CKAN 返回 success=false: {payload.get('error')}")
            break

        batch = payload["result"]["records"]
        records.extend(batch)
        print(f"[BACKFILL-NSW]   offset={offset} 拉到 {len(batch)} 条 (累计 {len(records)})")

        if len(batch) < PAGE_SIZE:
            break
        offset += PAGE_SIZE

    return records


# ══════════════════════════════════════════════════════════════
# 数据源 B：XLSX 下载（4 月份没有 datastore API）
# ══════════════════════════════════════════════════════════════

def fetch_xlsx_records(url):
    """下载 XLSX 并迭代成 dict 行。表头与 CSV/CKAN 一致。"""
    print(f"[BACKFILL-NSW] 下载 {url}")
    try:
        with urllib.request.urlopen(url, timeout=300) as resp:
            data = resp.read()
    except Exception as e:
        print(f"[BACKFILL-NSW] XLSX 下载失败: {e}")
        return []

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)

    try:
        header = [str(c).strip() if c is not None else "" for c in next(rows)]
    except StopIteration:
        return []

    records = []
    for row in rows:
        if row is None:
            continue
        records.append({header[i]: row[i] for i in range(len(header)) if i < len(row)})
    print(f"[BACKFILL-NSW]   XLSX 读到 {len(records)} 条")
    return records


# ══════════════════════════════════════════════════════════════
# 解析 + 聚合（与 WA 回填同样的 5 字段输出）
# ══════════════════════════════════════════════════════════════

def aggregate_records(records):
    """records 列含 FuelCode / Price / PriceUpdatedDate。返回 [{date,price,type,location,source}]"""
    buckets = defaultdict(list)   # (iso_date, type) -> [prices]

    for row in records:
        try:
            fuel_code = (row.get("FuelCode") or "").strip()
            fuel_type = FUEL_MAP.get(fuel_code)
            if fuel_type not in ALLOWED_TYPES:
                continue

            price_raw = row.get("Price")
            if price_raw in (None, ""):
                continue
            price = float(price_raw) / 100   # 分/升 → 澳元/升

            date_raw = str(row.get("PriceUpdatedDate") or "").strip()
            if not date_raw:
                continue
            # 兼容 "2026-02-01 02:18:03" 与 datetime 对象（XLSX 解析后可能直接是 datetime）
            day = date_raw[:10]
            try:
                datetime.strptime(day, "%Y-%m-%d")
            except ValueError:
                continue
            iso_date = f"{day}T00:00:00Z"

            buckets[(iso_date, fuel_type)].append(price)
        except Exception as e:
            print(f"[BACKFILL-NSW] 行解析错误: {e}")
            continue

    docs = []
    for (date, fuel_type), prices in buckets.items():
        docs.append({
            "date":     date,
            "price":    round(sum(prices) / len(prices), 4),
            "type":     fuel_type,
            "location": "NSW",
            "source":   "NSW FuelCheck",
        })
    return docs


# ══════════════════════════════════════════════════════════════
# 推入 Redis（与 fuel_harvester.push_to_redis 相同的去重 key）
# ══════════════════════════════════════════════════════════════

def push_to_redis(r, docs):
    pushed = skipped = 0
    for doc in docs:
        try:
            dedup_id = f"nsw:{doc['location']}:{doc['date'][:10]}:{doc['type']}"
            if r.sadd(DEDUP_KEY, dedup_id) == 0:
                skipped += 1
                continue
            r.lpush(QUEUE_KEY, json.dumps(doc, ensure_ascii=False))
            pushed += 1
        except Exception as e:
            print(f"[BACKFILL-NSW] Redis push 错误: {e}")
    return pushed, skipped


# ══════════════════════════════════════════════════════════════
# 主流程
# ══════════════════════════════════════════════════════════════

def main():
    r = redis.Redis(
        host=REDIS_HOST, port=REDIS_PORT,
        password=REDIS_PASS or None, decode_responses=True,
    )

    total_pushed = total_skipped = total_docs = 0

    # Feb / Mar 走 CKAN API
    for month, resource_id in DATASTORE_RESOURCES.items():
        print(f"[BACKFILL-NSW] === {month} (datastore_search) ===")
        records = fetch_datastore(resource_id)
        docs = aggregate_records(records)
        pushed, skipped = push_to_redis(r, docs)
        total_docs    += len(docs)
        total_pushed  += pushed
        total_skipped += skipped
        print(f"[BACKFILL-NSW] {month} -> docs={len(docs)} pushed={pushed} skipped={skipped}")

    # Apr 走 XLSX 下载
    for month, url in XLSX_URLS.items():
        print(f"[BACKFILL-NSW] === {month} (xlsx) ===")
        records = fetch_xlsx_records(url)
        docs = aggregate_records(records)
        pushed, skipped = push_to_redis(r, docs)
        total_docs    += len(docs)
        total_pushed  += pushed
        total_skipped += skipped
        print(f"[BACKFILL-NSW] {month} -> docs={len(docs)} pushed={pushed} skipped={skipped}")

    print(f"[BACKFILL-NSW] 完成 docs={total_docs} pushed={total_pushed} skipped={total_skipped}")


if __name__ == "__main__":
    main()
